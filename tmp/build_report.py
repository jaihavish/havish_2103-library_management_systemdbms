# -*- coding: utf-8 -*-
"""Fills the university PBL template with the library project content."""
import copy, json, os, re, sys
import docx
import mysql.connector
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(__file__))
import report_content as C

SRC = r"D:\denos\hp1\tmp\t.docx"
OUT = sys.argv[1] if len(sys.argv) > 1 else r"D:\denos\hp1\tmp\report_build.docx"
PAGES = json.load(open(r"D:\denos\hp1\tmp\pages.json")) if os.path.exists(r"D:\denos\hp1\tmp\pages.json") else {}
SHOTS = r"D:\denos\hp1\tmp\shots"
FONT = "Times New Roman"

d = docx.Document(SRC)
body = d.element.body
E = list(body)                    # original elements, by index (see dump_struct.py)
BODY = d._body


# ---------------------------------------------------------------- helpers
def style_run(r, size=None, bold=None, italic=None, font=FONT, color="000000", highlight=False):
    r.font.name = font
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    if size: r.font.size = Pt(size)
    if bold is not None: r.font.bold = bold
    if italic is not None: r.font.italic = italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if highlight: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def add_text(p, text, size=12, bold=None, italic=None, font=FONT):
    """Adds text; pieces wrapped in [[...]] become yellow placeholders."""
    for k, piece in enumerate(re.split(r"(\[\[.*?\]\])", text)):
        if not piece:
            continue
        ph = piece.startswith("[[")
        r = p.add_run(piece[2:-2].join("[]") if False else (("[" + piece[2:-2] + "]") if ph else piece))
        style_run(r, size, bold, italic, font, highlight=ph)


def remove(el):
    if el is not None and el.getparent() is not None:
        el.getparent().remove(el)


class Cursor:
    def __init__(self, el):
        self.cur = el

    def _put(self, el):
        self.cur.addnext(el)
        self.cur = el
        return Paragraph(el, BODY)

    def p(self, text="", size=12, align="j", bold=None, italic=None, space_after=6, line=1.5, keep_next=False,
          indent=None):
        P = self._put(OxmlElement("w:p"))
        pf = P.paragraph_format
        P.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER,
                       "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
        pf.line_spacing = line
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(0)
        if keep_next: pf.keep_with_next = True
        if indent is not None: pf.left_indent = Inches(indent)
        if text: add_text(P, text, size, bold, italic)
        return P

    def bullet(self, text, size=12):
        P = self._put(OxmlElement("w:p"))
        P.style = d.styles["List Bullet"]
        P.paragraph_format.line_spacing = 1.5
        P.paragraph_format.space_after = Pt(3)
        P.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_text(P, text, size)
        return P

    def labelled(self, label, text, size=12):
        P = self.p("", size)
        add_text(P, label, size, bold=True)
        add_text(P, text, size)
        return P

    def code(self, text, size=9.5):
        """Courier block, one paragraph per line, light grey background."""
        for line in text.rstrip("\n").split("\n"):
            P = self._put(OxmlElement("w:p"))
            pf = P.paragraph_format
            pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
            pf.left_indent = Inches(0.1)
            pf.keep_together = True
            pf.keep_with_next = True
            P.alignment = WD_ALIGN_PARAGRAPH.LEFT
            ppr = P._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
            ppr.append(shd)
            r = P.add_run(line if line else " ")
            style_run(r, size, font="Courier New")
        sp = self.p("", 6, space_after=4, line=1.0)
        return sp

    def picture(self, path, width_in):
        P = self._put(OxmlElement("w:p"))
        P.alignment = WD_ALIGN_PARAGRAPH.CENTER
        P.paragraph_format.keep_with_next = True
        P.paragraph_format.space_after = Pt(2)
        P.paragraph_format.line_spacing = 1.0
        P.add_run().add_picture(path, width=Inches(width_in))
        return P

    def caption(self, text, template_el):
        el = copy.deepcopy(template_el)
        for r in el.findall(qn("w:r"))[1:]:
            el.remove(r)
        self.cur.addnext(el)
        self.cur = el
        P = Paragraph(el, BODY)
        P.runs[0].text = text
        return P


def set_cell(cell, text, size=10, bold=False, align=None, font=FONT):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        remove(extra._p)
    for r in list(p.runs):
        remove(r._r)
    p.paragraph_format.space_after = Pt(0)
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        style_run(r, size, bold, None, font)
        if i < len(lines) - 1:
            r.add_break()
    if align:
        p.alignment = align
    tcpr = cell._tc.get_or_add_tcPr()
    va = tcpr.find(qn("w:vAlign"))
    if va is None:
        va = OxmlElement("w:vAlign"); tcpr.append(va)
    va.set(qn("w:val"), "top" if len(str(text)) > 28 else "center")


def fill_table(tbl_el, rows, widths=None, row_h=None, repeat_header=False, placeholder_cols=()):
    t = Table(tbl_el, BODY)
    trs = tbl_el.findall(qn("w:tr"))
    proto = trs[-1]
    for tr in trs[1:]:
        tbl_el.remove(tr)
    for row in rows:
        tr = copy.deepcopy(proto)
        tbl_el.append(tr)
    t = Table(tbl_el, BODY)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            set_cell(t.rows[ri].cells[ci], val)
        trpr = t.rows[ri]._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.insert(0, OxmlElement("w:cantSplit"))
        if row_h is not None:
            h = trpr.find(qn("w:trHeight"))
            if h is not None: h.set(qn("w:val"), str(row_h))
    if repeat_header:
        trpr = t.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
    if widths:
        grid = tbl_el.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for tr in tbl_el.findall(qn("w:tr")):
            for tc, w in zip(tr.findall(qn("w:tc")), widths):
                tcw = tc.get_or_add_tcPr().find(qn("w:tcW"))
                tcw.set(qn("w:w"), str(w)); tcw.set(qn("w:type"), "dxa")
    return t


def hl_cell(cell, text, size=10):
    """Set a cell to a yellow-highlighted placeholder."""
    set_cell(cell, text, size)
    for r in cell.paragraphs[0].runs:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def para_set_text(P, text, size=None):
    runs = P.runs
    runs[0].text = text
    for r in runs[1:]:
        if r._r.find(qn("w:br")) is None:
            remove(r._r)


def page_break_before(el):
    P = Paragraph(el, BODY)
    P.paragraph_format.page_break_before = True


def is_empty_par(el):
    if el.tag != qn("w:p"):
        return False
    if "".join(t.text or "" for t in el.iter(qn("w:t"))).strip():
        return False
    for tag in ("w:br", "w:drawing", "w:sectPr", "w:fldChar", "w:instrText"):
        if el.find(".//" + qn(tag)) is not None:
            return False
    return True


# ---------------------------------------------------------------- database facts
db = mysql.connector.connect(host="localhost", user="root", password=os.environ.get("LIB_DB_PASSWORD", "12345678"),
                             database="library_db")
cur = db.cursor(dictionary=True)
cur.execute("SELECT table_name AS t, column_name AS c, column_type AS ct, is_nullable AS n, column_key AS k, "
            "extra AS x, column_default AS df FROM information_schema.columns WHERE table_schema='library_db' "
            "AND table_name NOT LIKE 'v\\_%' ORDER BY table_name, ordinal_position")
COLS = cur.fetchall()
cur.execute("SELECT table_name AS t, column_name AS c, referenced_table_name AS rt FROM information_schema.key_column_usage "
            "WHERE table_schema='library_db' AND referenced_table_name IS NOT NULL")
FKS = {(r["t"], r["c"]): r["rt"] for r in cur.fetchall()}
TABLE_ORDER = [s[0] for s in C.SCHEMA]
counts = {}
for t in TABLE_ORDER:
    cur.execute(f"SELECT COUNT(*) AS n FROM `{t}`")
    counts[t] = cur.fetchone()["n"]
db.close()

DESC = {
    "publisher": {"publisher_id": "Unique publisher number", "name": "Publisher name", "city": "City of the publisher"},
    "author": {"author_id": "Unique author number", "name": "Author's full name", "country": "Author's country"},
    "category": {"category_id": "Unique category number", "name": "Category name"},
    "member_policy": {"member_type": "Member type: STUDENT, FACULTY or PUBLIC", "max_books": "Most books held at one time",
                      "loan_days": "Days a book may be kept", "fine_per_day": "Fine for each late day (Rs)"},
    "staff": {"staff_id": "Unique staff number", "username": "Login name", "full_name": "Staff member's name",
              "role": "LIBRARIAN or ADMIN"},
    "book": {"book_id": "Unique title number", "title": "Title of the book", "category_id": "Category of the title"},
    "book_author": {"book_id": "Title", "author_id": "Author of that title"},
    "edition": {"edition_id": "Unique edition number", "book_id": "Title this edition belongs to",
                "publisher_id": "Publisher of the edition", "isbn": "13-digit ISBN", "edition_no": "Edition number",
                "pub_year": "Year of publication", "pages": "Number of pages"},
    "copy": {"copy_id": "Unique copy number", "edition_id": "Edition this copy belongs to",
             "accession_no": "Library's own number for the copy", "shelf": "Shelf location",
             "status": "Where the copy is now (AVAILABLE, ISSUED, ON_HOLD, LOST, WITHDRAWN)",
             "added_on": "Date the copy was added"},
    "member": {"member_id": "Unique member number", "card_no": "Library card number", "full_name": "Member's name",
               "email": "E-mail address", "phone": "Phone number", "member_type": "Type; decides loan rules",
               "joined_on": "Date of joining", "expiry_date": "Date membership ends",
               "status": "ACTIVE, SUSPENDED or EXPIRED"},
    "issue": {"issue_id": "Unique loan number", "copy_id": "Copy lent", "member_id": "Borrower",
              "issued_by": "Staff member who issued it", "issue_date": "Date of issue", "due_date": "Date due back",
              "return_date": "Date returned; NULL while on loan",
              "active_copy_id": "Generated: copy_id while on loan, else NULL"},
    "reservation": {"reservation_id": "Unique reservation number", "edition_id": "Title reserved",
                    "member_id": "Member waiting", "queue_no": "Place in line; keeps counting up",
                    "reserved_on": "Date and time of reservation",
                    "status": "WAITING, READY, FULFILLED, CANCELLED or EXPIRED",
                    "copy_id": "Copy held for the member once READY", "ready_on": "Date the copy was set aside",
                    "live_queue_no": "Generated: queue_no while reservation is live",
                    "live_member_id": "Generated: member_id while reservation is live"},
    "fine": {"fine_id": "Unique fine number", "issue_id": "Late loan the fine belongs to",
             "days_late": "Days past the due date", "amount": "Fine amount in rupees",
             "paid_amount": "Amount paid so far", "assessed_on": "Date fine was raised",
             "paid_on": "Date of latest payment"},
    "digital_resource": {"resource_id": "Unique resource number", "title": "Title of the resource",
                         "resource_type": "EBOOK, JOURNAL, VIDEO or DATABASE", "publisher_id": "Supplier or publisher",
                         "url": "Address of the resource", "licence_end": "Date the licence ends"},
    "digital_access_log": {"log_id": "Unique log number", "resource_id": "Resource used", "member_id": "Member who used it",
                           "accessed_at": "Date and time of use", "minutes_used": "Minutes spent"},
}


def dict_rows():
    rows = []
    by = {}
    for c in COLS:
        by.setdefault(c["t"], []).append(c)
    for t in TABLE_ORDER:
        for i, c in enumerate(by[t]):
            ct = c["ct"].decode() if isinstance(c["ct"], bytes) else c["ct"]
            typ = "ENUM" if ct.startswith("enum") else ct.upper().replace("INT(", "INT(")
            x = c["x"].decode() if isinstance(c["x"], bytes) else c["x"]
            parts = []
            if c["k"] == "PRI": parts.append("PK")
            if (t, c["c"]) in FKS: parts.append("FK")
            if c["k"] == "UNI": parts.append("UNIQUE")
            if c["n"] == "NO" and c["k"] != "PRI" and "GENERATED" not in x: parts.append("NOT NULL")
            if "auto_increment" in x: parts.append("AUTO_INCREMENT")
            if "GENERATED" in x: parts.append("GENERATED")
            df = c["df"]
            if df is not None and "auto" not in x and "GENERATED" not in x:
                df = df.decode() if isinstance(df, bytes) else df
                parts.append("DEFAULT " + str(df).replace("DEFAULT_GENERATED", "").strip("()'") .replace("curdate", "today").replace("CURRENT_DATE", "today"))
            desc = DESC[t][c["c"]]
            if ct.startswith("enum"):
                desc += "" if "(" in desc else ""
            rows.append((t if i == 0 else "", c["c"], typ, ", ".join(parts), desc))
    return rows


# ---------------------------------------------------------------- 1. cover and front matter
set_cell_ph = hl_cell
t_title = Table(E[57], BODY)
set_cell(t_title.rows[0].cells[0], C.TITLE, 13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

t_info = Table(E[60], BODY)
hl_cell(t_info.rows[0].cells[1], C.PH.format("program, e.g. B.Tech CSE"), 10)
hl_cell(t_info.rows[1].cells[1], C.PH.format("section"), 10)
hl_cell(t_info.rows[2].cells[1], C.PH.format("faculty name"), 10)
hl_cell(t_info.rows[3].cells[1], C.PH.format("submission date"), 10)

t_by = Table(E[64], BODY)
hl_cell(t_by.rows[1].cells[0], C.PH.format("roll no."), 10)
set_cell(t_by.rows[1].cells[1], C.STUDENT_NAME, 10)
hl_cell(t_by.rows[1].cells[2], C.PH.format("mobile"), 10)
set_cell(t_by.rows[1].cells[3], C.STUDENT_EMAIL, 9)

# acknowledgement and undertaking: replace the two template placeholders inside runs
for el in (E[69], E[88]):
    P = Paragraph(el, BODY)
    for r in P.runs:
        if "Name of the Faculty" in r.text:
            r.text = r.text.replace("Name of the Faculty", "[Faculty Name]")
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if "[Project Title]" in r.text:
            r.text = r.text.replace("[Project Title]", C.TITLE)
t_sig = Table(E[93], BODY)
for row in t_sig.rows:
    for cell in row.cells:
        tx = cell.text.strip()
        if tx == "Name:":
            set_cell(cell, "Name: " + C.STUDENT_NAME, 11, bold=True)
        elif tx == "Roll Number:":
            hl_cell(cell, "Roll Number: " + C.PH.format("roll no."), 11)
        elif tx == "Program:":
            hl_cell(cell, "Program: " + C.PH.format("program"), 11)
for i in (86, 110, 144, 250, 331, 382):
    page_break_before(E[i])
page_break_before(E[301])

# acronyms table
ac = [a for a in C.ACRONYMS if a[0] not in ("CHECK", "DML", "GUI")]
fill_table(E[126], ac, widths=[2000, 7200])

# ---------------------------------------------------------------- abstract
cur_ = Cursor(E[144])
cur_.p(C.ABSTRACT, 12)
kw = Paragraph(E[151], BODY)
kw.runs[0].text = C.KEYWORDS
kw.runs[0].font.italic = False
kw.runs[0].font.size = Pt(12)
kw.paragraph_format.line_spacing = 1.5

# ---------------------------------------------------------------- chapter 1
Cursor(E[157]).p("", 1, space_after=0) if False else None
c = Cursor(E[157])
for para in C.PROBLEM:
    c.p(para)
c = Cursor(E[161])
for o in C.OBJECTIVES:
    c.bullet(o)
Cursor(E[168]).p("Table 1 states what the project covers and what it leaves out.", keep_next=True)
fill_table(E[170], C.SCOPE, widths=[2200, 7000])
Cursor(E[172]).p(C.STRATEGY_INTRO, keep_next=True)
fill_table(E[175], C.STRATEGY, widths=[2400, 6800])
Cursor(E[177]).p(C.FR_INTRO, keep_next=True)
fill_table(E[179], C.FR, widths=[900, 6900, 1400])
Cursor(E[181]).p(C.ERD_INTRO)
pic_par = Paragraph(E[183], BODY)
for r in list(pic_par.runs):
    remove(r._r)
pic_par.paragraph_format.keep_with_next = True
pic_par.add_run().add_picture(r"D:\denos\hp1\tmp\er.png", width=Inches(4.6))
Cursor(E[184]).p(C.ERD_NOTE)
Cursor(E[186]).p("Table 4 summarises the entities and their relationships.", keep_next=True)
fill_table(E[188], C.ENTITIES, widths=[2000, 3700, 3500], repeat_header=True)

# ---------------------------------------------------------------- chapter 2
Cursor(E[194]).p(C.SCHEMA_INTRO, keep_next=True)
fill_table(E[197], [(str(i + 1), n, p) for i, (n, p) in enumerate(C.SCHEMA)], widths=[800, 2300, 6100])
Cursor(E[199]).p(C.KEYS_INTRO, keep_next=True)
_names, _cur = [], ""
for _r in C.KEYS:
    _cur = _r[0] or _cur
    _names.append((_cur,) + tuple(_r[1:]))
fill_table(E[201], _names, widths=[1600, 2000, 3300, 2300], repeat_header=True)
_c = Cursor(E[201]); _c.p("", 4, space_after=0, line=1.0); _c.p(C.KEYS_NOTE)
c = Cursor(E[204])
c.p(C.FD_INTRO)
for fd in C.FDS:
    c.bullet(fd, 11)
Cursor(E[208]).p(C.NORM_INTRO, keep_next=True)
norm_rows = [(n, "Yes", "Yes", "Yes", "Yes", C.NORM_REMARKS[n]) for n in TABLE_ORDER]
fill_table(E[210], norm_rows, widths=[1700, 600, 600, 600, 700, 5000])
c = Cursor(E[210])
c.p("", 6, space_after=2, line=1.0)
for lab, txt in C.NORM_STEPS:
    c.labelled(lab, txt)
c.p(C.NORM_NOTE)
Cursor(E[212]).p("Table 8 is the data dictionary. It was generated from the live database, so the data types and keys "
                 "match what MySQL actually holds.", keep_next=True)
fill_table(E[215], dict_rows(), widths=[1450, 1600, 1250, 1750, 3150], row_h=300, repeat_header=True)

fill_table(E[220], C.ENV_DB, widths=[2300, 6900])
Cursor(E[217]).p(C.DDL_INTRO)
schema_sql = open(r"D:\denos\hp1\sql\01_schema.sql", encoding="utf8").read()


def ddl(name):
    m = re.search(r"CREATE TABLE " + name + r" \(.*?\n\);", schema_sql, flags=re.S)
    txt = m.group(0)
    return "\n".join(re.sub(r"\s*--.*$", "", l) for l in txt.split("\n") if not l.strip().startswith("--"))


c = Cursor(E[220])
c.p("The statements below show the three most important tables. The remaining tables follow the same style.", keep_next=True)
for nm in ("issue", "reservation", "fine"):
    c.code(ddl(nm), 8)
c.p(C.SAMPLE_INTRO, keep_next=True)
cnt = "\n".join(f"{t:<20}{counts[t]:>5}" for t in TABLE_ORDER)
c.code("table                rows\n--------------------------\n" + cnt)
c.p("Sample rows from the member and copy tables:", keep_next=True)
c.code("card_no     full_name        member_type  expiry_date  status\n"
       "LIB2025001  Aarav Sharma     STUDENT      +300 days    ACTIVE\n"
       "LIB2024101  Dr. Neha Kapoor  FACULTY      +800 days    ACTIVE\n"
       "LIB2025201  Manoj Kulkarni   PUBLIC       +165 days    ACTIVE\n"
       "LIB2023015  Ishaan Joshi     STUDENT      -30 days     EXPIRED\n\n"
       "accession_no  edition  shelf  status\n"
       "ACC-0001      1        CS-A1  ISSUED\n"
       "ACC-0002      1        CS-A1  AVAILABLE\n"
       "ACC-0029      13       HI-E2  ON_HOLD")

# ---------------------------------------------------------------- chapter 3
Cursor(E[254]).p(C.APP_OVERVIEW)
Cursor(E[258]).p(C.CONNECT_TEXT, keep_next=True)
fill_table(E[260], C.ENV_APP, widths=[2700, 6500])
Cursor(E[262]).p(C.MODULES_INTRO, keep_next=True)
fill_table(E[264], C.MODULES, widths=[1500, 2500, 3300, 1900], repeat_header=True)
Cursor(E[265]).p(C.CRUD_INTRO, keep_next=True)
fill_table(E[267], C.CRUD, widths=[1500, 3200, 2800, 1700])
Cursor(E[269]).p(C.RULES_INTRO, keep_next=True)
fill_table(E[272], C.RULES, widths=[2500, 3500, 3200], repeat_header=True)
Cursor(E[274]).p(C.REPORTS_INTRO, keep_next=True)
fill_table(E[276], C.REPORTS_TABLE, widths=[2200, 3800, 3200], repeat_header=True)

fig_cap_tpl = E[184]
c = Cursor(E[277])
c.p(C.SCREENS_INTRO)
for i, (fname, cap, lead) in enumerate(C.SCREENS, start=2):
    c.p(lead, keep_next=True)
    c.picture(os.path.join(SHOTS, fname + ".png"), 5.9)
    c.caption(f"Figure {i}. {cap}", fig_cap_tpl)

Paragraph(E[288], BODY).paragraph_format.keep_with_next = True
Cursor(E[288]).p(C.TEST_INTRO, keep_next=True)
fill_table(E[290], C.TESTS, widths=[850, 1400, 1950, 1700, 2200, 1100], repeat_header=True)
tt = Table(E[290], BODY)
for r in tt.rows[1:]:
    set_cell(r.cells[5], "Pass", 10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
c = Cursor(E[290])
c.p("", 6, space_after=2, line=1.0)
for para in C.TEST_RESULT:
    c.p(para)

# ---------------------------------------------------------------- chapter 4
c = Cursor(E[302])
c.p("", 6, space_after=2, line=1.0)
for para in C.CONCLUSION:
    c.p(para)
c = Cursor(E[307])
for f in C.FUTURE:
    c.bullet(f)

# ---------------------------------------------------------------- references
ref_tpl = E[333]
prev = ref_tpl
for i, txt in enumerate(C.REFS):
    el = ref_tpl if i == 0 else copy.deepcopy(ref_tpl)
    if i:
        prev.addnext(el)
    prev = el
    P = Paragraph(el, BODY)
    for r in P.runs[1:]:
        remove(r._r)
    P.runs[0].text = txt
    P.runs[0].font.size = Pt(11)
    P.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    P.paragraph_format.space_after = Pt(6)

# ---------------------------------------------------------------- appendices
c = Cursor(E[384])
c.p("The complete database script is submitted as four files in the sql folder of the project:", keep_next=True)
for f, desc in C.APPX_A:
    c.bullet(f"{f}: {desc}", 11)
c.p("Run them in the order shown. The first script drops and rebuilds the database named library_db.")
c = Cursor(E[388])
c.p("The application source code and the automated tests are in the project folder submitted with this report:", keep_next=True)
for f, desc in C.APPX_B:
    c.bullet(f"{f}: {desc}", 11)
c.p("To run: load the SQL files, then start the program with python app/app.py from the project folder.")
c = Cursor(E[392])
c.p(C.APPX_C_TEXT, keep_next=True)
views_sql = open(r"D:\denos\hp1\sql\03_views.sql", encoding="utf8").read()


def view(name):
    m = re.search(r"-- [^\n]*\nCREATE OR REPLACE VIEW " + name + r" AS.*?;\n", views_sql, flags=re.S)
    return m.group(0).strip()


c.code(view("v_overdue_copies"), 8.5)
c.code(view("v_reservation_queue"), 8.5)

# ---------------------------------------------------------------- clean the template
delete = set(range(0, 48)) | {67, 71, 72, 73, 74, 75, 76, 77, 78, 79, 80, 81, 82, 83, 84, 85, 87, 111, 112, 114,
                              116, 117, 120, 121, 124, 125, 145, 146, 147, 148, 150, 154, 155, 158, 159, 162, 163, 164,
                              165, 166, 173, 182, 193, 195, 202, 205, 206, 213, 218, 221, 253, 255, 256, 270, 278,
                              279, 280, 281, 282, 283, 284, 285, 286, 287, 291, 292, 304, 305, 308, 309, 310, 311,
                              385, 386, 389, 390, 393, 394}
delete |= set(range(334, 365))
for i in sorted(delete):
    remove(E[i])
keep_blank = {49, 50, 54, 55, 58, 59, 61, 62, 91, 92}
for i, el in enumerate(E):
    if i in keep_blank or el.getparent() is None:
        continue
    if is_empty_par(el):
        remove(el)

# ---------------------------------------------------------------- contents, lists, page numbers
toc = Table(E[113], BODY)
for row in toc.rows[1:]:
    key = row.cells[0].text.strip()
    pg = PAGES.get("toc", {}).get(key)
    if pg:
        set_cell(row.cells[1], pg, 11, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in row.cells[0].paragraphs[0].runs:
        r.font.size = Pt(11)

tbl_titles = []
fig_titles = []
for el in body.iter(qn("w:p")):
    txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
    m = re.match(r"(Table|Figure) (\d+)\. (.*)", txt)
    if m and Paragraph(el, BODY).style.name == "Caption":
        (tbl_titles if m.group(1) == "Table" else fig_titles).append((int(m.group(2)), m.group(3)))
tbl_titles.sort(); fig_titles.sort()
fill_table(E[118], [(f"Table {n}", t, PAGES.get("cap", {}).get(f"Table {n}", "")) for n, t in tbl_titles],
           widths=[1200, 6800, 1200], row_h=300)
fill_table(E[122], [(f"Figure {n}", t, PAGES.get("cap", {}).get(f"Figure {n}", "")) for n, t in fig_titles],
           widths=[1200, 6800, 1200], row_h=300)
for tb in (E[118], E[122]):
    for row in Table(tb, BODY).rows[1:]:
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

for el in body.iter(qn("w:p")):
    Pk = Paragraph(el, BODY)
    if Pk.style.name == "Caption" and re.match(r"Table \d+\.", Pk.text):
        Pk.paragraph_format.keep_with_next = True
for tb in body.iter(qn("w:tbl")):
    for pp in tb.findall(qn("w:tr"))[0].iter(qn("w:p")):
        Paragraph(pp, BODY).paragraph_format.keep_with_next = True
d.core_properties.title = "DBMS PBL Report - " + C.TITLE
d.core_properties.author = C.STUDENT_NAME
d.save(OUT)
print("saved", OUT)
