import json, re, sys
import pymupdf
import docx
pdf = pymupdf.open(sys.argv[1])
lines = [[l.strip() for l in p.get_text().split("\n")] for p in pdf]


def first(pred, start=0):
    for i in range(start, len(lines)):
        if any(pred(l) for l in lines[i]):
            return i + 1
    return ""


d = docx.Document(r"D:\denos\hp1\tmp\report_build.docx")
toc_tbl = d.tables[[t.rows[0].cells[0].text.strip() for t in d.tables].index("Contents")]
toc = {}
for row in toc_tbl.rows[1:]:
    key = row.cells[0].text
    k = key.strip()
    if k == "Abstract":
        pg = first(lambda l: l == "ABSTRACT", 7)
    elif k.startswith("LIST OF"):
        pg = first(lambda l, k=k: l == k, 4)
    elif re.match(r"^[1-4]\. ", k) and k[0] != "5":
        n = k[0]
        pg = first(lambda l, n=n: l == f"CHAPTER {n}", 8)
    elif k.startswith("5."):
        pg = first(lambda l: l == "REFERENCES", 8)
    elif k.startswith("6."):
        pg = first(lambda l: l == "APPENDICES", 8)
    else:
        num = k.split()[0]
        pg = first(lambda l, num=num: l.startswith(num + " ") and len(l) < 70, 8)
    toc[key] = str(pg)
# keys in the document table keep their leading spaces, build script strips them
toc = {k.strip(): v for k, v in toc.items()}
cap = {}
for kind, rng in (("Table", 15), ("Figure", 11)):
    for n in range(1, rng + 1):
        cap[f"{kind} {n}"] = str(first(lambda l, kind=kind, n=n: re.match(rf"^{kind} {n}\. ", l) is not None, 8))
json.dump({"toc": toc, "cap": cap}, open(r"D:\denos\hp1\tmp\pages.json", "w"), indent=1)
print(toc); print(cap)
