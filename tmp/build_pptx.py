import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Paths
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TMP_DIR = os.path.join(BASE_DIR, "tmp")
SHOTS_DIR = os.path.join(TMP_DIR, "shots")
ER_IMG = os.path.join(TMP_DIR, "er.png")

# Palette
DARK_BG = RGBColor(15, 23, 42)       # #0f172a (Deep Slate)
DARK_CARD = RGBColor(30, 41, 59)     # #1e293b
LIGHT_BG = RGBColor(248, 250, 252)   # #f8fafc
WHITE = RGBColor(255, 255, 255)
TEXT_DARK = RGBColor(30, 41, 59)
TEXT_MUTED = RGBColor(100, 116, 139)
TEXT_LIGHT = RGBColor(241, 245, 249)
PRIMARY_INDIGO = RGBColor(79, 70, 229) # #4f46e5
ACCENT_BLUE = RGBColor(37, 99, 235)    # #2563eb
ACCENT_GREEN = RGBColor(16, 185, 129)  # #10b981
ACCENT_ROSE = RGBColor(239, 68, 68)    # #ef4444
ACCENT_AMBER = RGBColor(245, 158, 11)  # #f59e0b
BORDER_COLOR = RGBColor(226, 232, 240) # #e2e8f0

def create_base_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_header(slide, title, category="DBMS PBL REVIEW", dark=False):
    # Category tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = PRIMARY_INDIGO if not dark else RGBColor(165, 180, 252)

    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_DARK if not dark else WHITE

def add_footer(slide, current_rev="REVIEW 2", dark=False):
    f_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
    tf = f_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"Library Circulation & Digital Resource Management System | DBMS PBL | {current_rev}"
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_MUTED if not dark else RGBColor(148, 163, 184)

def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_COLOR):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def add_bullet_list(tf, items, font_size=13, space_after=8, text_color=TEXT_DARK):
    for idx, item in enumerate(items):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.space_after = Pt(space_after)
        if isinstance(item, tuple):
            title, desc = item
            run1 = p.add_run()
            run1.text = "• " + title + ": "
            run1.font.bold = True
            run1.font.size = Pt(font_size)
            run1.font.color.rgb = text_color
            run2 = p.add_run()
            run2.text = desc
            run2.font.bold = False
            run2.font.size = Pt(font_size)
            run2.font.color.rgb = text_color
        else:
            p.text = "• " + str(item)
            p.font.size = Pt(font_size)
            p.font.color.rgb = text_color

# ==============================================================================
# REVIEW 2 PRESENTATION
# ==============================================================================
def generate_review_2():
    prs = create_base_deck()
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Dark Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = add_card(s1, Inches(0), Inches(0), Inches(13.333), Inches(7.5), bg_color=DARK_BG, border_color=None)
    
    # Title box
    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.0))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    
    p_badge = tf1.paragraphs[0]
    p_badge.text = "DATABASE MANAGEMENT SYSTEMS • PROJECT REVIEW 2"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(165, 180, 252)
    p_badge.space_after = Pt(14)
    
    p_title = tf1.add_paragraph()
    p_title.text = "DESIGN AND IMPLEMENTATION OF A DATABASE MANAGEMENT SYSTEM FOR LIBRARY CIRCULATION AND DIGITAL RESOURCE MANAGEMENT"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_after = Pt(12)
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "Focus: Conceptual E-R Modeling, Relational Schema Conversion, Normalization (1NF–3NF), Integrity Constraints & Schema DDL"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)
    p_sub.space_after = Pt(28)
    
    p_meta = tf1.add_paragraph()
    p_meta.text = "Student: Shreyash Singh  |  Technology: MySQL 8.0 & Python 3  |  Milestone: Review 2"
    p_meta.font.size = Pt(12)
    p_meta.font.bold = True
    p_meta.font.color.rgb = ACCENT_GREEN

    # Slide 2: Progress from Review 1 to Review 2
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Project Evolution: Progress from Review 1 to Review 2")
    add_footer(s2, "REVIEW 2")
    
    # Left Box: Review 1 Baseline
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    tb_r1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.6))
    tf_r1 = tb_r1.text_frame
    tf_r1.word_wrap = True
    p_h1 = tf_r1.paragraphs[0]
    p_h1.text = "Review 1 Scope & Problem Definition"
    p_h1.font.size = Pt(16)
    p_h1.font.bold = True
    p_h1.font.color.rgb = PRIMARY_INDIGO
    p_h1.space_after = Pt(12)
    add_bullet_list(tf_r1, [
        ("Problem Statement", "Manual registers and isolated spreadsheets cause duplicate data, untracked book copies, lost fines, and zero visibility into queue waiting lists."),
        ("Objectives Set", "Design an integrated, centralized relational database to manage print books, digital resources, patrons, circulation, and fines."),
        ("User Roles Identified", "Librarians (daily operations), Admins (policies & reports), and Members (students, faculty, public)."),
        ("Initial Scope", "High-level requirements gathering and initial feature specification.")
    ], font_size=12, space_after=10)

    # Right Box: Review 2 Deliverables Achieved
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_r2 = s2.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_r2 = tb_r2.text_frame
    tf_r2.word_wrap = True
    p_h2 = tf_r2.paragraphs[0]
    p_h2.text = "Review 2 Deliverables Achieved"
    p_h2.font.size = Pt(16)
    p_h2.font.bold = True
    p_h2.font.color.rgb = ACCENT_GREEN
    p_h2.space_after = Pt(12)
    add_bullet_list(tf_r2, [
        ("Comprehensive E-R Model", "Identified 15 distinct entities, cardinalities, associative entities, and relationship attributes."),
        ("3-Tier Catalog Architecture", "Split Book (title work), Edition (ISBN publication), and Copy (physical barcode/shelf item)."),
        ("Relational Mapping", "Transformed conceptual ER model into normalized relational schemas with full foreign key constraints."),
        ("Formal Normalization Proof", "Verified progression through 1NF, 2NF, and 3NF/BCNF to prevent update, insertion, and deletion anomalies."),
        ("MySQL 8.0 DDL Scripts", "Implemented 01_schema.sql with primary keys, checks, enums, cascades, and indexes.")
    ], font_size=12, space_after=10)

    # Slide 3: Functional Requirements & Business Rules
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "System Requirements & Circulation Rules Specification")
    add_footer(s3, "REVIEW 2")
    
    # 3 Column Cards
    col_w = Inches(3.7)
    gap = Inches(0.3)
    
    # Col 1: Functional Requirements
    add_card(s3, Inches(0.8), Inches(1.6), col_w, Inches(5.0))
    tb_c1 = s3.shapes.add_textbox(Inches(0.95), Inches(1.75), Inches(3.4), Inches(4.7))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "1. Functional Scope"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_c1, [
        ("Catalog Control", "Manage books, editions, authors, publishers, and categories."),
        ("Copy Tracking", "Unique accession numbers per physical copy with real-time status."),
        ("Circulation", "Issue, return, and loan renewal operations."),
        ("Hold Queue", "Ordered waitlist reservations for out-of-stock titles."),
        ("Fines & Audits", "Overdue fee calculation and payment tracking."),
        ("Digital Assets", "E-books, journals, and patron access duration logging.")
    ], font_size=11, space_after=6)

    # Col 2: Multi-Tier Loan Policies
    add_card(s3, Inches(0.8) + col_w + gap, Inches(1.6), col_w, Inches(5.0))
    tb_c2 = s3.shapes.add_textbox(Inches(0.95) + col_w + gap, Inches(1.75), Inches(3.4), Inches(4.7))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "2. Tiered Policy Table"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)
    add_bullet_list(tf_c2, [
        ("Student Policy", "Max 3 books concurrently | 14-day loan period | ₹2.00 / late day fine."),
        ("Faculty Policy", "Max 6 books concurrently | 30-day loan period | ₹1.00 / late day fine."),
        ("Public Policy", "Max 2 books concurrently | 7-day loan period | ₹5.00 / late day fine."),
        ("Separation Rationale", "Rules live in `member_policy` table so admins adjust loan rules dynamically without modifying codebase.")
    ], font_size=11, space_after=8)

    # Col 3: Critical Business Rules
    add_card(s3, Inches(0.8) + (col_w + gap)*2, Inches(1.6), col_w, Inches(5.0))
    tb_c3 = s3.shapes.add_textbox(Inches(0.95) + (col_w + gap)*2, Inches(1.75), Inches(3.4), Inches(4.7))
    tf_c3 = tb_c3.text_frame
    tf_c3.word_wrap = True
    p = tf_c3.paragraphs[0]
    p.text = "3. Integrity Constraints"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    p.space_after = Pt(10)
    add_bullet_list(tf_c3, [
        ("₹50 Fine Lock Rule", "A patron with > ₹50 in unpaid fines cannot borrow any new books."),
        ("1 Active Issue / Copy", "A physical copy cannot be issued if status is already ISSUED or ON_HOLD."),
        ("Reservation Trigger", "A reservation can only be placed when 100% of physical copies are checked out."),
        ("FIFO Hold Transition", "Returned books automatically assign to next WAITING reservation and shift to READY.")
    ], font_size=11, space_after=8)

    # Slide 4: Conceptual ER Model
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Conceptual Design: Entity-Relationship (ER) Architecture")
    add_footer(s4, "REVIEW 2")
    
    # Left: Details
    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.0))
    tb_er = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.6))
    tf_er = tb_er.text_frame
    tf_er.word_wrap = True
    p = tf_er.paragraphs[0]
    p.text = "Core Entities & Relationship Mappings"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_er, [
        ("Publisher (1) ──< (N) Edition", "A publisher produces many published book editions."),
        ("Book (N) >──< (M) Author", "Many-to-many relationship resolved using `book_author` bridge table."),
        ("Book (1) ──< (N) Edition", "One master title work has multiple physical published editions."),
        ("Edition (1) ──< (N) Copy", "One edition has multiple physical inventory copies on library shelves."),
        ("Copy (1) ──< (N) Issue", "A copy appears in loan history; only one active issue at a time."),
        ("Member (1) ──< (N) Issue", "Patrons borrow books governed by their linked `member_policy`."),
        ("Issue (1) ──< (0..1) Fine", "Late returned books generate at most one fine record."),
        ("Edition (1) ──< (N) Reservation", "Waitlist entries queue members for checked-out editions.")
    ], font_size=11, space_after=6)

    # Right: ER Diagram / Visual Representation
    add_card(s4, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.0))
    if os.path.exists(ER_IMG):
        s4.shapes.add_picture(ER_IMG, Inches(7.1), Inches(1.8), width=Inches(5.2))
    else:
        tb_diag = s4.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.2), Inches(4.0))
        tb_diag.text_frame.text = "ER Diagram Visual\n\nPUBLISHER ──< EDITION >── BOOK >── CATEGORY\n                │          │\n                │          └──< BOOK_AUTHOR >── AUTHOR\n                ^ N\n              COPY ──< ISSUE >── MEMBER >── MEMBER_POLICY\n                         │\n                       FINE"

    # Slide 5: The 3-Tier Catalog Architecture
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Key Structural Decision: The 3-Tier Catalog Model")
    add_footer(s5, "REVIEW 2")

    w3 = Inches(3.7)
    g3 = Inches(0.3)
    # Tier 1
    add_card(s5, Inches(0.8), Inches(1.6), w3, Inches(5.0))
    tb_t1 = s5.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(3.4), Inches(4.6))
    tf_t1 = tb_t1.text_frame
    tf_t1.word_wrap = True
    p = tf_t1.paragraphs[0]
    p.text = "Tier 1: BOOK (The Work)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_t1, [
        ("Concept", "Represents the intellectual literary work regardless of publication format."),
        ("Attributes", "`book_id` (PK), `title`, `category_id` (FK)."),
        ("Example", "'Database System Concepts'"),
        ("Advantage", "Title name, authors, and classification are defined exactly once. No duplication when multiple editions are purchased.")
    ], font_size=12, space_after=10)

    # Tier 2
    add_card(s5, Inches(0.8) + w3 + g3, Inches(1.6), w3, Inches(5.0))
    tb_t2 = s5.shapes.add_textbox(Inches(0.95) + w3 + g3, Inches(1.8), Inches(3.4), Inches(4.6))
    tf_t2 = tb_t2.text_frame
    tf_t2.word_wrap = True
    p = tf_t2.paragraphs[0]
    p.text = "Tier 2: EDITION (Publication)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)
    add_bullet_list(tf_t2, [
        ("Concept", "Represents a specific release from a publisher with page count and year."),
        ("Attributes", "`edition_id` (PK), `book_id` (FK), `publisher_id` (FK), `isbn` (UNIQUE 13-digit), `edition_no`, `pub_year`, `pages`."),
        ("Example", "7th Edition, McGraw-Hill, 2019, ISBN: 9780078022159."),
        ("Advantage", "Strictly separates publication metadata from physical items.")
    ], font_size=12, space_after=10)

    # Tier 3
    add_card(s5, Inches(0.8) + (w3 + g3)*2, Inches(1.6), w3, Inches(5.0))
    tb_t3 = s5.shapes.add_textbox(Inches(0.95) + (w3 + g3)*2, Inches(1.8), Inches(3.4), Inches(4.6))
    tf_t3 = tb_t3.text_frame
    tf_t3.word_wrap = True
    p = tf_t3.paragraphs[0]
    p.text = "Tier 3: COPY (Inventory)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(10)
    add_bullet_list(tf_t3, [
        ("Concept", "Represents a specific physical book sitting on the library shelf."),
        ("Attributes", "`copy_id` (PK), `edition_id` (FK), `accession_no` (UNIQUE barcode), `shelf`, `status` (AVAILABLE, ISSUED, ON_HOLD, WITHDRAWN)."),
        ("Example", "Barcode `ACC-CS-001` on shelf `CS-A1-01`."),
        ("Advantage", "Enables accurate per-item circulation history and real-time availability tracking.")
    ], font_size=12, space_after=10)

    # Slide 6: Normalization Proof (1NF to 3NF)
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Database Normalization: 1NF to 3NF / BCNF Analysis")
    add_footer(s6, "REVIEW 2")

    add_card(s6, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    tb_norm = s6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.6))
    tf_norm = tb_norm.text_frame
    tf_norm.word_wrap = True
    p = tf_norm.paragraphs[0]
    p.text = "Step-by-Step Normalization Breakdown"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(14)
    add_bullet_list(tf_norm, [
        ("First Normal Form (1NF) Compliance", "All attributes contain atomic values. Multi-valued fields (e.g. multiple authors for a book) are eliminated by introducing the associative table `book_author(book_id, author_id)`. ISBNs, phone numbers, and dates are stored in singular standardized columns."),
        ("Second Normal Form (2NF) Compliance", "All non-prime attributes are fully functionally dependent on the primary key. In composite primary key tables (like `book_author`), no non-key attribute depends on just a subset of the primary key. Author details live in `author`, title details live in `book`."),
        ("Third Normal Form (3NF) Compliance", "Elimination of transitive functional dependencies ($X \\rightarrow Y, Y \\rightarrow Z$). Example: A member's borrow limit and daily fine rate depend strictly on `member_type`, not `member_id`. Moving `(member_type, max_books, loan_days, fine_per_day)` into `member_policy` ensures that policy changes do not cause update anomalies across thousands of member rows."),
        ("Boyce-Codd Normal Form (BCNF)", "Every determinant in every relation is a candidate key. All functional dependencies $X \\rightarrow Y$ satisfy the condition that $X$ is a superkey.")
    ], font_size=12, space_after=12)

    # Slide 7: Database Relational Schema & Constraints
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Relational Schema Specification & Integrity Constraints")
    add_footer(s7, "REVIEW 2")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_s1 = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_s1 = tb_s1.text_frame
    tf_s1.word_wrap = True
    p = tf_s1.paragraphs[0]
    p.text = "Relational Table Catalog (15 Tables)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_s1, [
        ("Master Entities", "`publisher`, `author`, `category`, `staff`, `member_policy`"),
        ("Catalog Schema", "`book`, `book_author`, `edition`, `copy`"),
        ("Circulation Schema", "`member`, `issue`, `fine`, `reservation`"),
        ("Digital Resource Schema", "`digital_resource`, `digital_access_log`"),
        ("Primary Key Policy", "Auto-increment synthetic integers for internal joins; natural UNIQUE candidate keys for domain identifiers (`isbn`, `accession_no`, `card_no`).")
    ], font_size=12, space_after=10)

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_s2 = s7.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_s2 = tb_s2.text_frame
    tf_s2.word_wrap = True
    p = tf_s2.paragraphs[0]
    p.text = "Enforced Database Constraints"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    p.space_after = Pt(10)
    add_bullet_list(tf_s2, [
        ("Domain Constraints", "CHECK (`max_books` > 0), CHECK (`loan_days` > 0), CHECK (`fine_per_day` >= 0), CHECK (`days_late` >= 0)."),
        ("Enum Constraints", "`role` ENUM('LIBRARIAN', 'ADMIN'); `status` ENUM('AVAILABLE', 'ISSUED', 'ON_HOLD', 'LOST', 'WITHDRAWN'); `member_type` ENUM('STUDENT', 'FACULTY', 'PUBLIC')."),
        ("Referential Integrity", "Foreign key cascades on book deletion (`ON DELETE CASCADE`), while restricting deletion of copies currently out on loan (`RESTRICT`)."),
        ("Uniqueness", "Enforced unique ISBN-13, barcode accession numbers, member card numbers, and staff usernames.")
    ], font_size=12, space_after=10)

    # Slide 8: DDL Implementation (MySQL 8.0)
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Database Implementation: Schema DDL (`01_schema.sql`)")
    add_footer(s8, "REVIEW 2")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    tb_ddl = s8.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.6))
    tf_ddl = tb_ddl.text_frame
    tf_ddl.word_wrap = True
    p = tf_ddl.paragraphs[0]
    p.text = "DDL Implementation Highlights in MySQL 8.0"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(12)
    add_bullet_list(tf_ddl, [
        ("Engine & Character Set", "Configured all tables using MySQL `InnoDB` engine for robust ACID transaction support, foreign key constraints, and crash recovery. Default charset `utf8mb4` with `utf8mb4_unicode_ci` collation."),
        ("CHECK Constraint Enforcement", "Leveraged MySQL 8.0.16+ native CHECK constraint enforcement to ensure monetary fine amounts, days overdue, and book allocations never accept negative or invalid values."),
        ("Temporal Tracking", "`issue_date`, `due_date`, `return_date`, and `assessed_on` typed as DATE/DATETIME with DEFAULT `CURRENT_TIMESTAMP` for reliable circulation auditing."),
        ("Indexing Strategy", "Primary keys indexed automatically; added secondary indexes on `member.card_no`, `copy.accession_no`, `edition.isbn`, and `issue.return_date` to accelerate high-frequency circulation queries.")
    ], font_size=12, space_after=10)

    # Slide 9: Review 2 Summary & Roadmap to Review 3
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Summary of Review 2 & Roadmap to Final Review 3")
    add_footer(s9, "REVIEW 2")

    # Left: Completed in Review 2
    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_m1 = s9.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_m1 = tb_m1.text_frame
    tf_m1.word_wrap = True
    p = tf_m1.paragraphs[0]
    p.text = "Review 2 Deliverables Complete"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)
    add_bullet_list(tf_m1, [
        ("Conceptual ER Model", "Complete entity identification & relationship cardinality defined."),
        ("3-Tier Normalization", "Separated Book, Edition, Copy; normalized to 3NF/BCNF."),
        ("Relational Schema", "Full 15-table relational design with primary and foreign keys."),
        ("Integrity Architecture", "CHECK constraints, ENUM validations, and uniqueness rules."),
        ("MySQL DDL Generation", "`01_schema.sql` scripted, verified, and executed in MySQL 8.0.")
    ], font_size=12, space_after=10)

    # Right: Plan for Review 3
    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_m2 = s9.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_m2 = tb_m2.text_frame
    tf_m2.word_wrap = True
    p = tf_m2.paragraphs[0]
    p.text = "Target Roadmap for Review 3"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(12)
    add_bullet_list(tf_m2, [
        ("Data Seeding", "Populate database with comprehensive realistic test data (`02_sample_data.sql`)."),
        ("SQL Views & Queries", "Create operational views and analytical queries (joins, aggregations, window functions)."),
        ("Transaction Management", "Implement atomic transaction engine with pessimistic locking in Python (`library.py`)."),
        ("Dual User Interface", "Build full desktop GUI (Tkinter) and modern Web Application (Flask + Tailwind)."),
        ("Systematic Testing", "Complete unit and integration test suite (`test_library.py`).")
    ], font_size=12, space_after=10)

    rev2_path = os.path.join(BASE_DIR, "Review_2_Presentation.pptx")
    prs.save(rev2_path)
    print("Saved Review 2 presentation to:", rev2_path)
    return rev2_path

# ==============================================================================
# REVIEW 3 PRESENTATION
# ==============================================================================
def generate_review_3():
    prs = create_base_deck()
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Dark Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = add_card(s1, Inches(0), Inches(0), Inches(13.333), Inches(7.5), bg_color=DARK_BG, border_color=None)
    
    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.0))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    
    p_badge = tf1.paragraphs[0]
    p_badge.text = "DATABASE MANAGEMENT SYSTEMS • FINAL PROJECT REVIEW 3"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(165, 180, 252)
    p_badge.space_after = Pt(14)
    
    p_title = tf1.add_paragraph()
    p_title.text = "DESIGN AND IMPLEMENTATION OF A DATABASE MANAGEMENT SYSTEM FOR LIBRARY CIRCULATION AND DIGITAL RESOURCE MANAGEMENT"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_after = Pt(12)
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "Final Phase: Advanced SQL Queries, Analytical Views, ACID Transactions, Dual Front-End (Web & Desktop), Systematic Testing & Deployment"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)
    p_sub.space_after = Pt(28)
    
    p_meta = tf1.add_paragraph()
    p_meta.text = "Student: Shreyash Singh  |  Technology: MySQL 8.0, Python 3, Flask, Tkinter  |  Milestone: Final Review 3"
    p_meta.font.size = Pt(12)
    p_meta.font.bold = True
    p_meta.font.color.rgb = ACCENT_GREEN

    # Slide 2: Progress from Review 2 to Review 3
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Implementation Journey: Progress from Review 2 to Review 3", category="DBMS PBL REVIEW 3")
    add_footer(s2, "REVIEW 3")

    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0))
    tb_r2 = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.6))
    tf_r2 = tb_r2.text_frame
    tf_r2.word_wrap = True
    p = tf_r2.paragraphs[0]
    p.text = "State at Review 2 (Design Phase)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(12)
    add_bullet_list(tf_r2, [
        ("Conceptual ER Model", "Entity-Relationship diagram completed with relationships and cardinalities."),
        ("Normalized Relational Schema", "15 tables normalized up to 3NF/BCNF."),
        ("DDL Scripts Created", "`01_schema.sql` written with primary, foreign keys and check constraints."),
        ("Limitation", "Database was empty; no circulation queries, views, transactions, or user applications existed yet.")
    ], font_size=12, space_after=10)

    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_r3 = s2.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_r3 = tb_r3.text_frame
    tf_r3.word_wrap = True
    p = tf_r3.paragraphs[0]
    p.text = "Implemented in Review 3 (Final Phase)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)
    add_bullet_list(tf_r3, [
        ("Comprehensive Data Population", "`02_sample_data.sql` with realistic books, members, issues, fines, and digital logs."),
        ("Pre-compiled SQL Views", "6 operational views created (`03_views.sql`) including window functions."),
        ("Complex SQL Queries", "10 benchmark queries (`04_queries.sql`) covering multi-joins, aggregates, and subqueries."),
        ("ACID Transaction Engine", "Python business layer (`library.py`) with atomic transactions and pessimistic locking."),
        ("Dual Interface Applications", "Built both a desktop GUI (`app.py`) and a modern web application (`web_app.py`)."),
        ("Systematic Unit Testing", "Automated test suite (`test_library.py`) verifying all business rules.")
    ], font_size=12, space_after=8)

    # Slide 3: Sample Data Seeding & Database Population
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Database Population & Realism (`02_sample_data.sql`)", category="DBMS PBL REVIEW 3")
    add_footer(s3, "REVIEW 3")

    add_card(s3, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    tb_data = s3.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.6))
    tf_data = tb_data.text_frame
    tf_data.word_wrap = True
    p = tf_data.paragraphs[0]
    p.text = "Realistic Seed Dataset Overview"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(12)
    add_bullet_list(tf_data, [
        ("Diverse Book Catalogue", "Populated with real-world computer science, mathematics, physics, and literature textbooks (e.g. Silberschatz Database Concepts, Cormen Algorithms, Rosen Discrete Mathematics) across diverse publishers (MIT Press, McGraw-Hill, Wiley, Prentice Hall)."),
        ("Multi-Tier Member Base", "Registered Student, Faculty, and Public patrons with valid card numbers, emails, and phone numbers."),
        ("Varied Circulation States", "Created active loans on shelf copies, on-time returned books, and realistic overdue loans to evaluate penalty calculations."),
        ("Fine Accrual Records", "Populated both settled and outstanding fine records with exact days overdue and fine assessments."),
        ("Hold Reservation Queues", "Simulated high-demand titles with patrons waiting in ordered queue positions (`WAITING` and `READY`)."),
        ("Digital Resource Logs", "Recorded patron access sessions on online learning platforms (ACM Digital Library, O'Reilly, MIT OCW).")
    ], font_size=12, space_after=10)

    # Slide 4: Advanced SQL Queries & Analytical Demonstrations
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Advanced SQL Query Implementation (`04_queries.sql`)", category="DBMS PBL REVIEW 3")
    add_footer(s4, "REVIEW 3")

    w2 = Inches(5.7)
    add_card(s4, Inches(0.8), Inches(1.6), w2, Inches(5.0))
    tb_q1 = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_q1 = tb_q1.text_frame
    tf_q1.word_wrap = True
    p = tf_q1.paragraphs[0]
    p.text = "Joins & Aggregate Evaluations"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_q1, [
        ("Multi-Table Relational Joins", "5-table join combining `book`, `edition`, `publisher`, `book_author`, and `author` with `GROUP_CONCAT` to aggregate multi-author names."),
        ("Aggregations with GROUP BY & HAVING", "Query Q8 groups issues by member and uses `HAVING COUNT(*) >= 3` to isolate frequent borrowers."),
        ("Category Volume Aggregates", "Query Q9 groups categories with `LEFT JOIN` and computes distinct titles and total physical copies.")
    ], font_size=12, space_after=12)

    add_card(s4, Inches(6.8), Inches(1.6), w2, Inches(5.0))
    tb_q2 = s4.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_q2 = tb_q2.text_frame
    tf_q2.word_wrap = True
    p = tf_q2.paragraphs[0]
    p.text = "Nested Queries & Subqueries"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)
    add_bullet_list(tf_q2, [
        ("Nested NOT IN Subquery", "Query Q10 isolates dead inventory—books that have never been issued—by nesting a subquery selecting issued copy IDs."),
        ("Dynamic Search Subquery", "Query Q3 powers catalog search matching titles, author names, categories, or ISBNs simultaneously."),
        ("Fine Accrual Calculations", "Calculates live fines by combining `DATEDIFF(CURDATE(), i.due_date)` with dynamic member policy rates.")
    ], font_size=12, space_after=12)

    # Slide 5: SQL Views Architecture
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Analytical Database Views (`03_views.sql`)", category="DBMS PBL REVIEW 3")
    add_footer(s5, "REVIEW 3")

    add_card(s5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    tb_v = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.6))
    tf_v = tb_v.text_frame
    tf_v.word_wrap = True
    p = tf_v.paragraphs[0]
    p.text = "Pre-Compiled Views for Data Abstraction & Real-Time Reporting"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(12)
    add_bullet_list(tf_v, [
        ("v_available_titles", "Real-time list of titles with at least one physical copy on the shelf (`status = 'AVAILABLE'`), with distinct authors concatenated."),
        ("v_overdue_copies", "Joins `issue`, `copy`, and `member` with `DATEDIFF` to compute exact overdue days and accrued fines per borrower."),
        ("v_member_history", "Complete historical loan audit per member showing issue date, due date, return date, loan status, and paid fines."),
        ("v_popular_books", "Aggregates circulation counts by book to highlight high-demand textbooks and reading materials."),
        ("v_reservation_queue", "Uses the SQL Window Function `ROW_NUMBER() OVER (PARTITION BY edition_id ORDER BY queue_no)` to present live queue positions."),
        ("v_fine_summary", "Monthly revenue reporting aggregating fines assessed, total collected, and still pending.")
    ], font_size=12, space_after=10)

    # Slide 6: Transaction Management & ACID Guarantees
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Transaction Management & Concurrency Control", category="DBMS PBL REVIEW 3")
    add_footer(s6, "REVIEW 3")

    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_t1 = s6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_t1 = tb_t1.text_frame
    tf_t1.word_wrap = True
    p = tf_t1.paragraphs[0]
    p.text = "ACID Properties Implementation"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_t1, [
        ("Atomicity", "All multi-table mutations (e.g. issuing a book updates `copy.status`, inserts `issue`, and updates `reservation`) run inside a Python transaction context manager (`tx()`). If any step errors, the entire block rolls back."),
        ("Consistency", "Foreign keys and CHECK constraints prevent corrupt states; borrow limit and fine lock validations precede any write."),
        ("Isolation", "Transactions prevent dirty and non-repeatable reads using MySQL InnoDB default REPEATABLE READ isolation."),
        ("Durability", "Committed transactions persist to disk via MySQL write-ahead redo logs.")
    ], font_size=12, space_after=8)

    add_card(s6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_t2 = s6.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_t2 = tb_t2.text_frame
    tf_t2.word_wrap = True
    p = tf_t2.paragraphs[0]
    p.text = "Pessimistic Concurrency Locking"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    p.space_after = Pt(10)
    add_bullet_list(tf_t2, [
        ("SELECT ... FOR UPDATE", "Used during issue, return, and reservation operations. Locks the target `copy` and `member` rows to prevent race conditions when two librarians attempt to issue the same book simultaneously."),
        ("Queue Number Serialization", "Locks the edition reservation rows to guarantee sequential, un-duplicated `queue_no` assignment during concurrent reservation requests."),
        ("Automatic Rollbacks", "Handles integrity errors and custom `LibraryError` exceptions gracefully, reporting clean feedback to users.")
    ], font_size=12, space_after=8)

    # Slide 7: Application Layer & Circulation Engine
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Business Logic & Circulation Engine (`library.py`)", category="DBMS PBL REVIEW 3")
    add_footer(s7, "REVIEW 3")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    tb_be = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.6))
    tf_be = tb_be.text_frame
    tf_be.word_wrap = True
    p = tf_be.paragraphs[0]
    p.text = "Python Business Layer Architecture"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(12)
    add_bullet_list(tf_be, [
        ("Shared Codebase", "`app/library.py` contains all circulation logic, shared identically across the desktop GUI (`app.py`), the web app (`web_app.py`), and the automated test suite (`tests/test_library.py`)."),
        ("Input Validation Engine", "Custom regex sanitization: `check_isbn` verifies strict 13-digit format; `check_email` verifies valid email syntax; `check_phone` enforces 10 digits."),
        ("₹50 Fine Lock Enforcement", "Queries outstanding unpaid fines for the borrower; blocks book checkouts if debt exceeds ₹50.00."),
        ("Book Return & Dynamic Fine Calculation", "Computes `(today - due_date)` in days and multiplies by `fine_per_day` for that borrower's tier; inserts a `fine` record only if overdue."),
        ("FIFO Hold Queue Transition", "When a book is returned, `_free_copy()` checks for WAITING reservations. If found, the copy is reserved as `ON_HOLD` and reservation status becomes `READY`.")
    ], font_size=12, space_after=10)

    # Slide 8: User Interfaces - Desktop GUI & Web Application
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Front-End Applications: Desktop GUI & Web Interface", category="DBMS PBL REVIEW 3")
    add_footer(s8, "REVIEW 3")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_u1 = s8.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_u1 = tb_u1.text_frame
    tf_u1.word_wrap = True
    p = tf_u1.paragraphs[0]
    p.text = "1. Desktop GUI (`app/app.py`)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_u1, [
        ("Framework", "Python Tkinter with custom `ttk` styling (`clam` engine)."),
        ("Design System", "Deep navy header, card-based layout, Segoe UI typography."),
        ("Key Screens", "Dashboard overview, Catalogue search & copy adder, Circulation issue/return, Member directory, Reservations queue, Fines ledger, Reports exporter."),
        ("Best For", "Librarian desktop workstations with local peripheral barcode scanners.")
    ], font_size=12, space_after=10)

    add_card(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_u2 = s8.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_u2 = tb_u2.text_frame
    tf_u2.word_wrap = True
    p = tf_u2.paragraphs[0]
    p.text = "2. Modern Web App (`web_app.py`)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(10)
    add_bullet_list(tf_u2, [
        ("Stack", "Python Flask REST API + HTML5 / Tailwind CSS + Chart.js."),
        ("Design System", "Modern SaaS responsive layout with real-time donut and bar charts."),
        ("Interactive Features", "Instant live search, animated toast alerts, interactive issue/return dropdowns, member history drawer, 1-click fine payments, and CSV export."),
        ("Zero-Friction Fallback", "Includes live MySQL connection modal AND zero-setup in-memory fallback demo mode.")
    ], font_size=12, space_after=8)

    # Slide 9: System Demonstration & Screenshots
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Application Demonstration: Core Operational Screens", category="DBMS PBL REVIEW 3")
    add_footer(s9, "REVIEW 3")

    # Add 2 screenshot cards side by side
    dash_img = os.path.join(SHOTS_DIR, "dashboard.png")
    circ_img = os.path.join(SHOTS_DIR, "issue.png")
    
    w_shot = Inches(5.7)
    add_card(s9, Inches(0.8), Inches(1.6), w_shot, Inches(5.0))
    tb_d1 = s9.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.3), Inches(0.6))
    tb_d1.text_frame.text = "Executive Dashboard & Overdue Board"
    tb_d1.text_frame.paragraphs[0].font.bold = True
    tb_d1.text_frame.paragraphs[0].font.color.rgb = PRIMARY_INDIGO
    if os.path.exists(dash_img):
        s9.shapes.add_picture(dash_img, Inches(1.0), Inches(2.3), width=Inches(5.3))

    add_card(s9, Inches(6.8), Inches(1.6), w_shot, Inches(5.0))
    tb_d2 = s9.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.3), Inches(0.6))
    tb_d2.text_frame.text = "Circulation Desk (Book Issue & Return)"
    tb_d2.text_frame.paragraphs[0].font.bold = True
    tb_d2.text_frame.paragraphs[0].font.color.rgb = ACCENT_GREEN
    if os.path.exists(circ_img):
        s9.shapes.add_picture(circ_img, Inches(7.0), Inches(2.3), width=Inches(5.3))

    # Slide 10: Systematic Unit Testing & Verification
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Systematic Testing & Rule Verification (`tests/test_library.py`)", category="DBMS PBL REVIEW 3")
    add_footer(s10, "REVIEW 3")

    add_card(s10, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    tb_test = s10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.6))
    tf_test = tb_test.text_frame
    tf_test.word_wrap = True
    p = tf_test.paragraphs[0]
    p.text = "Automated Test Suite Methodology & Test Scenarios"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(12)
    add_bullet_list(tf_test, [
        ("Isolated Test Harness", "Each test run drops and rebuilds `library_db` from `01_schema.sql`, `02_sample_data.sql`, and `03_views.sql` to guarantee repeatability."),
        ("Borrow Limit Enforcement Tests", "Verifies that student patrons holding 3 books are strictly blocked upon attempting to check out a 4th book."),
        ("Fine Lock Verification Tests", "Verifies that an unpaid fine exceeding ₹50.00 blocks borrowing, while a partial payment bringing debt below ₹50 restores borrowing rights."),
        ("Overdue Math Validation Tests", "Simulates overdue dates and asserts that calculated fines exactly equal `(days_late * fine_per_day)` according to member tier."),
        ("Reservation Queue Fulfillment Tests", "Tests that returning a reserved book puts the copy into `ON_HOLD`, advances the reservation to `READY`, and prevents other patrons from taking it."),
        ("Transaction Rollback Tests", "Forces database errors and verifies that partial writes leave zero residual records in `copy` or `issue`.")
    ], font_size=12, space_after=10)

    # Slide 11: Technical Achievements & Key Learnings
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Technical Achievements & Project Outcomes", category="DBMS PBL REVIEW 3")
    add_footer(s11, "REVIEW 3")

    col_w = Inches(3.7)
    gap = Inches(0.3)

    add_card(s11, Inches(0.8), Inches(1.6), col_w, Inches(5.0))
    tb_o1 = s11.shapes.add_textbox(Inches(0.95), Inches(1.75), Inches(3.4), Inches(4.7))
    tf_o1 = tb_o1.text_frame
    tf_o1.word_wrap = True
    p = tf_o1.paragraphs[0]
    p.text = "1. Data Integrity"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_o1, [
        ("Normalized 3NF", "Zero data redundancy across titles, authors, editions, and copies."),
        ("Integrity Constraints", "Strict CHECK and foreign key rules prevent invalid entries at the database engine level."),
        ("ACID Guarantees", "Atomic transactions protect multi-table updates from corruption.")
    ], font_size=11, space_after=8)

    add_card(s11, Inches(0.8) + col_w + gap, Inches(1.6), col_w, Inches(5.0))
    tb_o2 = s11.shapes.add_textbox(Inches(0.95) + col_w + gap, Inches(1.75), Inches(3.4), Inches(4.7))
    tf_o2 = tb_o2.text_frame
    tf_o2.word_wrap = True
    p = tf_o2.paragraphs[0]
    p.text = "2. Operational Performance"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)
    add_bullet_list(tf_o2, [
        ("Pre-Compiled Views", "Eliminates repetitive complex join queries for common circulation screens."),
        ("Indexing Architecture", "Primary and secondary B-Tree indexes speed up lookups on barcodes and ISBNs."),
        ("Pessimistic Locking", "Prevents concurrency collisions during peak circulation periods.")
    ], font_size=11, space_after=8)

    add_card(s11, Inches(0.8) + (col_w + gap)*2, Inches(1.6), col_w, Inches(5.0))
    tb_o3 = s11.shapes.add_textbox(Inches(0.95) + (col_w + gap)*2, Inches(1.75), Inches(3.4), Inches(4.7))
    tf_o3 = tb_o3.text_frame
    tf_o3.word_wrap = True
    p = tf_o3.paragraphs[0]
    p.text = "3. Usability & Deliverables"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(10)
    add_bullet_list(tf_o3, [
        ("Dual Interface", "Provides both desktop GUI and responsive browser web interface."),
        ("Automated Tests", "Comprehensive test suite verifying business logic."),
        ("Complete Documentation", "Formal project reports in Markdown, Word, PDF, and PPTX.")
    ], font_size=11, space_after=8)

    # Slide 12: Conclusion & Future Scope
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Conclusion & Future Enhancements", category="DBMS PBL REVIEW 3")
    add_footer(s12, "REVIEW 3")

    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_c1 = s12.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "Project Summary"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_INDIGO
    p.space_after = Pt(10)
    add_bullet_list(tf_c1, [
        ("Complete Lifecycle", "Successfully executed the entire database development lifecycle from requirement analysis, conceptual modeling, normalization, physical implementation, to full-stack application deployment."),
        ("Objective Achieved", "Replaced uncoordinated registers with an integrated, reliable DBMS enforcing circulation rules, queue order, fine settlement, and digital usage tracking."),
        ("Ready for Production", "Equipped with dual user interfaces, comprehensive test suite, and database views.")
    ], font_size=12, space_after=10)

    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    tb_c2 = s12.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "Future Enhancements"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)
    add_bullet_list(tf_c2, [
        ("Hardware Integration", "Direct integration with handheld Barcode / QR / RFID scanners for automated shelf audits."),
        ("Automated Notifications", "Automated email / SMS webhooks notifying patrons 2 days before due dates and alerting reservation holders when books become READY."),
        ("Online Student Portal", "Self-service mobile web portal allowing students to view active loans and renew books remotely.")
    ], font_size=12, space_after=10)

    rev3_path = os.path.join(BASE_DIR, "Review_3_Presentation.pptx")
    prs.save(rev3_path)
    print("Saved Review 3 presentation to:", rev3_path)
    return rev3_path

if __name__ == "__main__":
    generate_review_2()
    generate_review_3()
